from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = "111111"
TEXT = "242424"
MUTED = "767676"
LIGHT = "E3E3E3"
PALE = "F5F5F5"
WHITE = "FFFFFF"
FONT = "Malgun Gothic"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), BLACK)
        borders.append(node)
    for edge in ("insideH",):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "2")
        node.set(qn("w:color"), LIGHT)
        borders.append(node)
    for edge in ("left", "right", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def set_font(run, size: float, *, bold=False, color=TEXT, italic=False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str, size=10.5, color=TEXT) -> None:
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size - 0.3, bold=True, color=BLACK)
        else:
            run = paragraph.add_run(part)
            set_font(run, size, color=color)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # LibreOffice can clip the leading digit of two-digit PAGE fields when a
    # right-aligned footer sits exactly on the printable edge. Keep a small
    # inset so pages 10 and above render their full number.
    paragraph.paragraph_format.right_indent = Inches(0.12)
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, separate, text, end):
        run._r.append(node)
    set_font(run, 8, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.78)
    # Keep flowing body text clear of the footer in both Word and LibreOffice.
    # The prior 0.72-inch margin let long pages share the page-number baseline.
    section.bottom_margin = Inches(0.98)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = "AI PRONOTE v1.5 사용설명서  |  ㈜써니엔터테인먼트"
    for run in footer_p.runs:
        set_font(run, 8, color=MUTED)
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_field(footer.add_paragraph())


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("AI PRONOTE  PRODUCT MANUAL")
    set_font(run, 8.5, bold=True, color=MUTED)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(36)
    p_pr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "22")
    bottom.set(qn("w:color"), BLACK)
    borders.append(bottom)
    p_pr.append(borders)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("AI PRONOTE")
    set_font(run, 40, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    add_inline(subtitle, "노트와 녹음 그리고 AI 회의록 사용설명서", 20, BLACK)

    desc = doc.add_paragraph()
    desc.paragraph_format.space_after = Pt(28)
    add_inline(desc, "설치부터 음성·영상 녹화, 필기, 받아쓰기, AI 회의록, 데이터 복구까지 설명합니다.", 11.5, MUTED)

    meta = [
        ("제작", "㈜써니엔터테인먼트"),
        ("제품", "AI PRONOTE v1.5"),
        ("버전", "v1.5.0-beta13.20260920.2"),
        ("발행", "2026년 9월 20일"),
        ("대상", "Windows · macOS · iPad · 휴대폰 외부 테스트"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for i, (label, value) in enumerate(meta):
        table.columns[0].width = Inches(1.0)
        table.columns[1].width = Inches(5.5)
        for cell in table.rows[i].cells:
            set_cell_margins(cell, 90, 60, 90, 60)
        lp = table.cell(i, 0).paragraphs[0]
        vp = table.cell(i, 1).paragraphs[0]
        set_font(lp.add_run(label), 9, color=MUTED)
        set_font(vp.add_run(value), 10, bold=True, color=BLACK)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int) -> None:
    if level == 2:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(18)
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.keep_with_next = True
        match = re.match(r"(\d{2}\.?)\s*(.*)", text)
        if match:
            set_font(paragraph.add_run(match.group(1) + "  "), 13.5, bold=True, color="9A9A9A")
            set_font(paragraph.add_run(match.group(2)), 13.5, bold=True, color=BLACK)
        else:
            set_font(paragraph.add_run(text), 13.5, bold=True, color=BLACK)
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), LIGHT)
        bottom.set(qn("w:space"), "5")
        borders.append(bottom)
        p_pr.append(borders)
    else:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_with_next = True
        set_font(paragraph.add_run(text), 11.5, bold=True, color=BLACK)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_fill(cell, BLACK)
                color, bold = WHITE, True
            else:
                if r_idx % 2 == 0:
                    set_cell_fill(cell, PALE)
                color, bold = TEXT, False
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            add_inline(p, value.strip(), 8.8, color)
            for run in p.runs:
                run.font.bold = bold or run.font.bold
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(raw) > 1 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in raw[1]):
        raw.pop(1)
    return raw, i


def add_body_from_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    try:
        i = lines.index("---") + 1
    except ValueError:
        i = 1
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), 2)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), 3)
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), PALE)
            p._p.get_or_add_pPr().append(shading)
            add_inline(p, stripped[2:], 9.8, "303030")
            for run in p.runs:
                run.font.italic = True
            i += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.*)", stripped)
        bullet = re.match(r"^-\s+(.*)", stripped)
        if ordered or bullet:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(3)
            prefix = f"{ordered.group(1)}.  " if ordered else "—  "
            set_font(p.add_run(prefix), 10.2, bold=bool(ordered), color=BLACK)
            add_inline(p, ordered.group(2) if ordered else bullet.group(1), 10.2)
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped, 10.5)
        i += 1


def build(source: Path, output: Path) -> None:
    markdown = source.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_body_from_markdown(doc, markdown)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "AI PRONOTE v1.5 사용설명서"
    doc.core_properties.subject = "설치 녹음 영상 필기 받아쓰기 AI 회의록 사용법"
    doc.core_properties.author = "㈜써니엔터테인먼트"
    doc.core_properties.keywords = "AI PRONOTE, 사용설명서, 회의 녹음, 필기"
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_beta13_manual.py SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
